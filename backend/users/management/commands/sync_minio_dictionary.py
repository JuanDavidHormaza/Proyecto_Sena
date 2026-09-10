"""
===============================================================================
Comando: sync_minio_dictionary

Sincroniza UNA SOLA VEZ el diccionario de palabras (Excel subido a MinIO,
bucket "worklex-palabras") con las imágenes y audios que YA fueron subidos
manualmente a los buckets "worklex-images" / "worklex-audios".

Esto NO deja a la aplicación dependiendo del Excel: solo se usa como fuente
de datos para poblar la base de datos (DigitalDictionary + MediaAsset). Una
vez importado, la app funciona 100% desde PostgreSQL + MinIO, sin volver a
leer el archivo en cada request.

Uso:
    python manage.py sync_minio_dictionary
    python manage.py sync_minio_dictionary --program "Desarrollo de Software"
    python manage.py sync_minio_dictionary --excel-bucket worklex-palabras --excel-key "ADSO VOCABULARY.xlsx"
===============================================================================
"""

import io
import re
import unicodedata

from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from users.Models.modelsSENA import DigitalDictionary, MediaAsset, Subject
from users.services import media_storage


def _normalize_key(value: str) -> str:
    """'Use Case' -> 'usecase' (sin acentos, espacios ni mayúsculas)."""
    value = (value or "").strip()
    if not value:
        return ""
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.lower()
    value = re.sub(r"[\s_-]+", "", value)
    value = re.sub(r"[^a-z0-9]", "", value)
    return value


def _strip_extension(filename: str) -> str:
    return re.sub(r"\.[a-zA-Z0-9]+$", "", filename)


class Command(BaseCommand):
    help = "Sincroniza el diccionario (Excel en MinIO) con las imágenes/audios ya subidos."

    def add_arguments(self, parser):
        parser.add_argument(
            "--program",
            default="Desarrollo de Software",
            help="Programa SENA al que se asignan estas palabras (por defecto: Desarrollo de Software).",
        )
        parser.add_argument(
            "--ficha",
            default="",
            help="Ficha a la que se asignan (vacío = contenido general del programa, sin ficha específica).",
        )
        parser.add_argument(
            "--excel-bucket",
            default="worklex-palabras",
            help="Bucket de MinIO donde está el Excel de vocabulario.",
        )
        parser.add_argument(
            "--excel-key",
            default="ADSO VOCABULARY.xlsx",
            help="Nombre del archivo Excel dentro del bucket.",
        )
        parser.add_argument(
            "--subject",
            default="MULTIMEDIA",
            help="subject_id de la asignatura a la que se asocian las palabras.",
        )

    def handle(self, *args, **options):
        program = options["program"]
        ficha = options["ficha"]
        excel_bucket = options["excel_bucket"]
        excel_key = options["excel_key"]
        subject_id = options["subject"]

        subject = Subject.objects.filter(subject_id=subject_id).first()
        if not subject:
            self.stderr.write(self.style.ERROR(f"No existe la asignatura '{subject_id}'."))
            return

        self.stdout.write(f"Descargando '{excel_key}' desde bucket '{excel_bucket}'...")
        try:
            excel_bytes = media_storage.download_object_bytes(excel_bucket, excel_key)
        except Exception as exc:
            self.stderr.write(self.style.ERROR(f"No se pudo descargar el Excel: {exc}"))
            return

        workbook = load_workbook(io.BytesIO(excel_bytes), data_only=True)
        sheet = workbook.active
        self.stdout.write(f"Hoja leída: '{sheet.title}' ({sheet.max_row} filas)")

        # ── Indexar imágenes y audios ya existentes en MinIO ────────────────
        self.stdout.write("Indexando imágenes y audios en MinIO...")
        images = media_storage.list_objects(media_storage.BUCKET_IMAGES)
        audios = media_storage.list_objects(media_storage.BUCKET_AUDIOS)

        images_by_key = {}
        for item in images:
            filename = item["object_key"].rsplit("/", 1)[-1]
            key = _normalize_key(_strip_extension(filename))
            if key and key not in images_by_key:
                images_by_key[key] = item

        audios_by_key = {}
        for item in audios:
            filename = item["object_key"].rsplit("/", 1)[-1]
            key = _normalize_key(_strip_extension(filename))
            if key and key not in audios_by_key:
                audios_by_key[key] = item

        self.stdout.write(f"  {len(images_by_key)} imágenes indexadas, {len(audios_by_key)} audios indexados.")

        # ── Procesar filas del Excel (Word, Definition, Synonyms, ...) ──────
        created_words = 0
        updated_words = 0
        linked_images = 0
        linked_audios = 0
        media_assets_created = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            word = str(row[0]).strip()
            definition = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            synonyms = str(row[2]).strip() if len(row) > 2 and row[2] else ""

            # Columna "Difficulty" del Excel ADSO (escala 2-9). Es la que
            # decide en qué nivel MCER (A1-B2) aparece esta palabra.
            raw_difficulty = row[3] if len(row) > 3 else None
            try:
                difficulty = int(float(raw_difficulty)) if raw_difficulty is not None else 2
            except (ValueError, TypeError):
                difficulty = 2
            difficulty = max(2, min(9, difficulty))

            if not word:
                continue

            word_key = _normalize_key(word)
            image_match = images_by_key.get(word_key)
            audio_match = audios_by_key.get(word_key)

            image_url = image_match["url"] if image_match else ""
            audio_url = audio_match["url"] if audio_match else ""

            doc, was_created = DigitalDictionary.objects.update_or_create(
                word_id=word,
                subject=subject,
                defaults={
                    "definition": definition,
                    "synonyms": synonyms,
                    "image": image_url,
                    "audio": audio_url,
                    "video": "",
                    "program": program,
                    "ficha": ficha,
                    "difficulty": difficulty,
                },
            )
            if was_created:
                created_words += 1
            else:
                updated_words += 1

            # Registrar también en MediaAsset para que aparezcan en la
            # jerarquía Programa > Ficha > Tipo (y su inversa).
            if image_match and image_url:
                linked_images += 1
                _, made = MediaAsset.objects.get_or_create(
                    bucket=media_storage.BUCKET_IMAGES,
                    object_key=image_match["object_key"],
                    defaults={
                        "media_type": "image",
                        "program": program,
                        "ficha": ficha,
                        "word_id": word,
                        "definition": definition,
                        "synonyms": synonyms,
                        "subject": subject,
                        "url": image_url,
                        "original_filename": image_match["object_key"].rsplit("/", 1)[-1],
                        "size_bytes": image_match.get("size") or 0,
                    },
                )
                if made:
                    media_assets_created += 1

            if audio_match and audio_url:
                linked_audios += 1
                _, made = MediaAsset.objects.get_or_create(
                    bucket=media_storage.BUCKET_AUDIOS,
                    object_key=audio_match["object_key"],
                    defaults={
                        "media_type": "audio",
                        "program": program,
                        "ficha": ficha,
                        "word_id": word,
                        "definition": definition,
                        "synonyms": synonyms,
                        "subject": subject,
                        "url": audio_url,
                        "original_filename": audio_match["object_key"].rsplit("/", 1)[-1],
                        "size_bytes": audio_match.get("size") or 0,
                    },
                )
                if made:
                    media_assets_created += 1

        # Distribución de dificultad importada (para confirmar que los 4
        # niveles MCER van a tener palabras disponibles).
        from django.db.models import Count
        difficulty_counts = (
            DigitalDictionary.objects.filter(program=program)
            .values("difficulty")
            .annotate(total=Count("id"))
            .order_by("difficulty")
        )
        difficulty_summary = ", ".join(
            f"dif={row['difficulty']}: {row['total']}" for row in difficulty_counts
        )

        self.stdout.write(self.style.SUCCESS(
            f"\nListo. Palabras creadas: {created_words}, actualizadas: {updated_words}\n"
            f"Imágenes enlazadas: {linked_images}, audios enlazados: {linked_audios}\n"
            f"Registros MediaAsset nuevos: {media_assets_created}\n"
            f"Programa asignado: '{program}' | Ficha: '{ficha or '(sin ficha)'}'\n"
            f"Distribución de dificultad: {difficulty_summary}"
        ))
